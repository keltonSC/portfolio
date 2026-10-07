from io import BytesIO
from pathlib import Path
import os
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import pandas as pd

from export_safe import dataframe_to_excel_bytes


class ExportTests(unittest.TestCase):
    def test_untrusted_strings_cannot_become_excel_formulas(self):
        texts = ["=1+1", "+1", "-1", "@SUM(A1)", "\t=1+1", "  =1+1", "Texto fictício"]
        frame = pd.DataFrame({"Texto": texts, "Número": [10] * len(texts)})
        workbook = load_workbook(BytesIO(dataframe_to_excel_bytes(frame)), data_only=False)
        rows = list(workbook.active.iter_rows(min_row=2))
        for index, row in enumerate(rows):
            self.assertNotEqual(row[0].data_type, "f")
            self.assertEqual(row[1].value, 10)
            self.assertEqual(row[0].value, ("'" if index < 6 else "") + texts[index])
        self.assertEqual(frame.iloc[0]["Texto"], "=1+1")
        workbook.close()


class AnalyticsRegression(unittest.TestCase):
    def test_archived_lead_fallback_and_diagnostic_without_private_samples(self):
        import streamlit_app as app
        payload = {"data": [{"data_cad": "2026-01-02", "situacao": {"nome": "Arquivado"}, "empreendimento": [{"nome": "Empreendimento fictício"}], "corretor": None, "corretor_anterior": {"nome": "Responsável fictício"}, "momento": "Texto de teste"}]}
        state = {}
        with patch.object(app, "_leads_fetch_raw", return_value=(200, payload, "")), patch.object(app.st, "session_state", state):
            result = app.cv_list_period("demo", "service@example.invalid", "fixture-only", "2026-01-01", "2026-01-31")
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0]["Corretor Final"], "Responsável fictício")
        self.assertEqual(str(result.iloc[0]["Situacao (Target)"]), "Arquivado")
        self.assertNotIn("broker_samples", state["__leads_diag__"])

    def test_actual_session_adapter_never_retries(self):
        import streamlit_app as app
        with app.get_http_session() as session:
            self.assertEqual(session.get_adapter("https://").max_retries.total, 0)

    def test_same_service_query_is_not_shared_global_cache(self):
        import streamlit_app as app
        with patch.object(app, "_leads_fetch_raw", return_value=(200, {"data": []}, "")) as fetch, patch.object(app.st, "session_state", {}):
            for _ in range(2):
                app.cv_list_period_cached("demo", "service@example.invalid", "fixture-only", "2026-01-01", "2026-01-31", 20)
        self.assertEqual(fetch.call_count, 2)

    def test_browser_state_cannot_change_tenant_used_with_service_credentials(self):
        from streamlit.testing.v1 import AppTest
        with patch.dict(os.environ, {"RELATORIO_LIVE_ENABLED": "true"}), patch("security.require_authorized_user", return_value=None), patch("security.crm_json", return_value=(200, {"data": []})) as query:
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "streamlit_app.py"))
            app.secrets["cv_leads"] = {"subdomain": "configured-tenant", "email": "service@example.invalid", "token": "fixture-only"}
            app.secrets["cv_atendimento"] = dict(app.secrets["cv_leads"])
            app.session_state["cv_sub"] = "untrusted-tenant"
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertGreater(query.call_count, 0)
            self.assertEqual(query.call_args_list[0].args[1], "configured-tenant")


if __name__ == "__main__":
    unittest.main()
